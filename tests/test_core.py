"""独立预期顺序、原文件哈希、样式、异常与双格式回归。"""
import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path
from aligner.core import *
from aligner.export import export
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill

class AlignmentTests(unittest.TestCase):
    def pair(self):
        return detect(from_text('图档编号\nABC-2-R\nABC-1-R\nABC-1-X')),detect(from_text('零件编号\t数量\nABC-1\t2\nMISS\t1\nABC-2\t3\nABC-1\t4'))

    def test_direction_and_stability(self):
        r,s=self.pair();result=align(r,s)
        self.assertEqual([d['original_row'] for d in result['rows']],[4,2,5,3])
        self.assertEqual(result['ambiguous'],2)
        self.assertEqual(result['rows'][1]['reference_rows'],[3,4])
        self.assertEqual(align(detect(from_text('编号\nABC')),detect(from_text('编号\nABC-X')))['matched'],0)

    def test_fixed_policy(self):
        r,s=self.pair()
        self.assertEqual([d['original_row'] for d in align(r,s,'fixed')['rows']],[4,3,2,5])

    def test_empty_and_unicode(self):
        self.assertEqual(norm(' ＡＢＣ-００１ '),'ABC-001')
        with self.assertRaises(ValueError): from_text('  ')

    def test_style_and_integrity(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'in.xlsx';out=Path(td)/'out.xlsx'
            w=Workbook();s=w.active;s.append(['零件编号','数量']);s.append(['B-002',3]);s.append(['A-001',1])
            s['A2'].font=Font(color='AA0000',bold=True);s['A2'].fill=PatternFill('solid',fgColor='EEFFFF')
            s.column_dimensions['A'].width=38;s.row_dimensions[2].height=31
            w.create_sheet('其他工作表')['A1']='保留';w.save(p)
            before=hashlib.sha256(p.read_bytes()).hexdigest()
            src=choose(load_tables(p));res=align(detect(from_text('编号\nA-001\nB-002')),src);export(src,res,out)
            new=load_workbook(out)
            self.assertEqual(new.active['A3'].value,'B-002')
            self.assertEqual(new.active['A3'].font.color.rgb,'00AA0000')
            self.assertEqual(new.active.row_dimensions[3].height,31)
            self.assertEqual(new.active.column_dimensions['A'].width,38)
            self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),before)
            with zipfile.ZipFile(p) as a,zipfile.ZipFile(out) as b:
                for name in a.namelist():
                    if name!=src.part:self.assertEqual(a.read(name),b.read(name))
            with self.assertRaises(ValueError):export(src,res,p)

    def test_corrupt_file(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'坏文件.xlsx';p.write_bytes(b'bad')
            with self.assertRaisesRegex(ValueError,'无法读取'):load_tables(p)

    def test_reference_overwrite_and_formula_guard(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'in.xlsx';refpath=Path(td)/'ref.xlsx';out=Path(td)/'out.xlsx'
            w=Workbook();s=w.active;s.append(['零件编号','数量']);s.append(['B-002','=1+1']);s.append(['A-001',1]);w.save(p)
            refpath.write_bytes(b'keep reference')
            src=choose(load_tables(p));res=align(detect(from_text('编号\nA-001\nB-002')),src)
            with self.assertRaisesRegex(ValueError,'参考'):export(src,res,refpath,refpath)
            self.assertEqual(refpath.read_bytes(),b'keep reference')
            with self.assertRaisesRegex(ValueError,'公式'):export(src,res,out)
            self.assertFalse(out.exists())

    def test_headerless_and_footer(self):
        t=detect(from_text('Q-002\nQ-001'))
        self.assertEqual(t.header,0)
        t=detect(from_text('零件编号\t数量\nQ-002\t1\n客户码：测试\t制表：'))
        self.assertEqual(t.records(),[2])

    def test_xls_export(self):
        import xlwt,xlrd
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'in.xls';out=Path(td)/'out.xls';w=xlwt.Workbook();s=w.add_sheet('表')
            for i,v in enumerate(['零件编号','B-002','A-001']):s.write(i,0,v)
            s.col(0).width=8000;w.save(str(p))
            src=choose(load_tables(p));export(src,align(detect(from_text('编号\nA-001\nB-002')),src),out)
            result=xlrd.open_workbook(out,formatting_info=True).sheet_by_index(0)
            self.assertEqual(result.col_values(0),['零件编号','A-001','B-002'])
            self.assertEqual(result.colinfo_map[0].width,8000)

    def test_paste_formula_as_text(self):
        with tempfile.TemporaryDirectory() as td:
            r,s=self.pair();s.rows[2][2]='=1+1'
            p=Path(td)/'out.xlsx';export(s,align(r,s),p)
            w=load_workbook(p)
            self.assertEqual(w.active['B3'].data_type,'s')

if __name__=='__main__':unittest.main(verbosity=2)
