"""生成可公开复现的模拟业务样例，不包含附件企业资料。"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def styled(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = '模拟业务数据'
    ws.append(['模拟业务数据 仅用于功能演示'])
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1] + ws[2]:
        cell.font = Font(name='微软雅黑',bold=True,color='FFFFFF')
        cell.fill = PatternFill('solid',fgColor='174F55')
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.font = Font(name='微软雅黑',size=11)
            cell.fill = PatternFill('solid',fgColor='E8F3F0' if cell.row % 2 else 'FFFFFF')
            cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[row[0].row].height = 26
    for c in 'ABCDE':
        ws.column_dimensions[c].width = 25
    ws.freeze_panes = 'A3'
    path.parent.mkdir(parents=True,exist_ok=True)
    wb.save(path)

def main():
    base = Path(__file__).resolve().parents[1] / 'samples'
    styled(base/'模拟参考.xlsx',['图档编号','零件名称','数量'],[
        ['M-003-REV2','底板',1],['M-001-REV1','夹板',2],['M-002','垫片',3],
        ['M-004-A','支架甲',1],['M-004-B','支架乙',1],['M-005','压板',2],['M-006','导套',4]])
    styled(base/'模拟待排序.xlsx',['批次','零件编号','零件名称','数量','备注'],[
        ['演示','M-002','垫片',3,'精确匹配'],['演示','M-001','夹板',2,'部分匹配'],
        ['演示','M-099','试制件',1,'无参考编号'],['演示','M-003','底板',1,'部分匹配'],
        ['演示','M-004','支架',2,'一对多候选'],['演示','M-002','垫片补料',1,'重复编号稳定排序'],
        ['演示','M-006','导套',4,'精确匹配'],['演示','M-005','压板',2,'精确匹配']])
    for name in ['参考','待排序']:
        from aligner.core import load_tables
        t=load_tables(base/f'模拟{name}.xlsx')[0]
        rows=['\t'.join(str(t.value(r,c)) for c in sorted({c for v in t.rows.values() for c in v})) for r in sorted(t.rows) if r>=2]
        (base/f'模拟{name}.tsv').write_text('\n'.join(rows),encoding='utf-8')

if __name__=='__main__':
    main()
