"""从实际验证记录生成申报书；完整参赛信息只从本机私有文件读取。"""
from pathlib import Path
import json,argparse
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]

def font(run,size=11,bold=False):
    run.font.name='宋体';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)

def para(doc,text='',style=None):
    p=doc.add_paragraph(style=style);font(p.add_run(text))
    p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.35
    return p

def heading(doc,text):
    p=doc.add_paragraph(style='Heading 1');font(p.add_run(text),16,True)
    p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(12)

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for r,row in enumerate(t.rows):
        for c,w in zip(row.cells,widths):
            c.width=Cm(w)
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EDF0' if r==0 else ('F6F8FA' if r%2==0 else 'FFFFFF'));tcPr.append(shading)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','bottom','left','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.15
                for run in p.runs:font(run,10,r==0)
    para(doc)
    return t

def build(path,info):
    cases=json.loads((ROOT/'validation/attachments_summary.json').read_text(encoding='utf-8'))['cases']
    bench=json.loads((ROOT/'validation/benchmark.json').read_text())
    soak=json.loads((ROOT/'validation/soak.json').read_text())
    total=sum(c['total'] for c in cases)
    d=Document();s=d.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21);s.top_margin=Cm(2.2);s.bottom_margin=Cm(2.0);s.left_margin=Cm(2.3);s.right_margin=Cm(2.3)
    for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=d.styles[sty];st.font.name='宋体';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        for border in list(st._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    d.core_properties.author='项目申报人';d.core_properties.title='序衡报表智能排序对齐项目申报书';d.core_properties.subject='报表智能排序对齐算法';d.core_properties.version='1.0.0'
    p=d.add_paragraph(style='Title');font(p.add_run('序衡报表智能排序对齐项目申报书'),22,True);p.paragraph_format.space_after=Pt(20)
    para(d,'比赛题目  报表智能排序对齐算法')
    para(d,'作品版本  1.0.0     申报日期  2026年10月15日')
    table(d,['参赛信息','填写内容'],[['姓名',info['name']],['单位',info['organization']],['联系方式',info['contact']]], [3.1,13.3])
    heading(d,'项目简介')
    para(d,'序衡是一款面向报表整理工作的本地桌面应用。用户提供参考表格和待排序表格，程序自动找到编号列，按参考顺序重排完整数据行，再导出新的 Excel 文件和逐行审计记录。文件导入、文本粘贴、结果查看和成果导出已经在同一套程序中连通。')
    para(d,'我选择这个题目，是因为同一批零件资料常常来自不同部门。表头、列位置和编号写法不一致，手工复制排序既费心，也容易让编号和数量错行。我希望把顺序对齐这一步做得直接，同时让使用者能够查清每一行为什么被放到当前位置。')
    heading(d,'交付成果')
    para(d,'本项目已形成 Python 完整源码、Windows 桌面启动入口和打包程序、可复现模拟样例、自动化验证脚本、中文配音与同步字幕演示视频，以及本申报书。公开交付使用模拟资料，比赛附件中的原始业务表格仅在本机验证。')
    para(d,f'本轮实际处理了附件提供的 10 组输入，共 {total:,} 行待排序数据。另以固定随机种子的 1,000 行模拟表验证已知预期顺序，并保留每次运行记录。')
    d.add_page_break()
    heading(d,'一 使用场景与操作流程')
    para(d,'使用场景是将采购、模具或零件明细对齐到已有清单。参考表负责定义顺序；待排序表提供需要重排的完整业务行。即使两张表的列数、表头位置和工作表数量不同，程序也先识别编号，再建立关系。')
    table(d,['步骤','用户操作','程序输出'],[
        ['导入资料','分别选择 .xls 或 .xlsx 文件，或粘贴表格文本','读取有效单元格并预览来源顺序'],
        ['自动识别','点击导入并识别','显示所选工作表、编号列和表头行'],
        ['运行对齐','选择未匹配处理策略，点击开始对齐','展示匹配数量、排序结果和全部参考候选'],
        ['核对成果','查看原始行号、参考行号和异常颜色','区分已匹配、多候选与未匹配'],
        ['导出文件','填写本机保存位置，点击导出成果','生成新 Excel、审计 CSV 和 JSON']
    ],[2.1,7.0,7.3])
    heading(d,'处理规则')
    para(d,'匹配方向遵循题目示例：待排序编号必须包含在参考编号中。完全相同的编号优先；否则进行单向子串匹配。例如待排序项 ABC-001 可以关联 ABC-001-REV2，反过来不成立。归一化只处理全半角、大小写和首尾空白，保留内部字符。')
    para(d,'同一来源编号可以对应多个参考位置，全部候选会进入审计记录，排序使用最早参考位置；多个来源行对应同一参考项时，保留来源行的先后顺序，不复制或删除业务行。多候选数量属于已匹配数量的一部分。')
    para(d,'题目对未匹配行同时出现“保留原位置”和“置于末尾”两种描述。应用提供两种策略：默认将未匹配项放在末尾并保留它们的相对次序；也可固定未匹配行的位置，仅对其余位置中的匹配项排序。两种策略均记录原始行号。')
    d.add_page_break()
    heading(d,'二 技术实现与数据保护')
    para(d,'项目采用 Python 3.12 和 Tkinter 构建桌面界面，文件处理放在后台线程，结果刷新在界面线程完成。命令行与桌面版共用同一核心，便于复现和批量验证。运行时不调用网络服务，不需要模型或 API Key。')
    table(d,['模块','已经实现的职责'],[
        ['core.py','稀疏读取 XLSX、解析 XLS、粘贴解析、列评分、跨表工作表选择、匹配和稳定排序'],
        ['export.py','XLSX 行片段重排、XLS 值型表格导出、审计文件、原子写入和本地日志'],
        ['gui.py 与 main.py','桌面工作台、输入预览、结果着色、保存位置输入和命令行入口'],
        ['tests 与 scripts','独立预期顺序回归、附件验证、千行计时、循环运行和交付生成']
    ],[4.1,12.3])
    heading(d,'自动识别的依据')
    para(d,'列评分综合编号表头、非空比例、字符串长度、唯一性和字母数字形态。对于多工作表文件，程序进一步比较与待排序编号的覆盖关系，避免误选相似的备件页。XLSX 按实际存在的单元格读取，避免远端空白样式把扫描范围扩展到最后一列。')
    heading(d,'文件和日志')
    para(d,'XLSX 导出直接重排行片段，保留单元格样式引用、行高、列宽和标题区；其他 ZIP 部件保持字节一致。排序状态独立导出为审计表，不往原业务表中插入额外列。输出先写入临时文件，成功后再替换目标，禁止覆盖来源文件。')
    para(d,'本地操作日志按月存储在用户本机的 Xuheng 日志目录，记录时间、文件名和问题数量，不自动清理，执行不少于 180 天的保留策略。图纸内容、附件表格和编号清单不会发送到模型或语音服务。演示旁白单独使用中文神经语音合成，送出的只有公开解说文字。')
    para(d,'对损坏或加密文件给出错误提示。为保护数据，移动区域出现公式或跨行合并时，应用会提示先转换为值或拆分合并；XLS 导出适用于值型工作表。常规列宽、字体、颜色和 XLSX 非目标部件的保留已纳入回归检查。')
    d.add_page_break()
    heading(d,'三 实际运行与验证结果')
    para(d,'验证日期为 2026年10月15日，环境为 Windows 11、Python 3.12.14。附件只有输入样本，没有独立官方预期输出。下表记录程序实际运行得到的数量和完整导入导出时间，匹配数量不等同于官方标注下的匹配准确率。')
    table(d,['样本','待排序行','匹配行','未匹配','多候选','总耗时 秒'],[
        [str(c['case']).zfill(2),c['total'],c['matched'],c['unmatched'],c['ambiguous'],f"{c['end_to_end_s']:.3f}"] for c in cases
    ],[1.4,2.6,2.6,2.3,2.3,5.2])
    para(d,f'10 组输入共 {total:,} 行，输出逐行值与来源业务行保持一致，来源文件哈希均未改变。每组结果均通过独立规则基准的顺序比对；其他 XLSX 部件均保持字节一致。这里的顺序结论仅针对已声明的匹配和排序规则。')
    para(d,f'固定随机种子的 1,000 行模拟表独立预期顺序验证通过。5 次从导入到导出及审计写入的中位耗时为 {bench["median_seconds"]:.3f} 秒，范围为 {min(bench["end_to_end_seconds"]):.3f} 至 {max(bench["end_to_end_seconds"]):.3f} 秒。该结果描述本机本轮实测，不替代指定 i7 环境下的验收。')
    para(d,f'循环运行实测持续 {soak["actual_seconds"]:.2f} 秒，完成 {soak["iterations"]:,} 次样例导入、匹配和导出，异常数为 {soak["errors"]}。交付包含可设为 72 小时的耐久脚本；本轮短时记录不作 72 小时稳定性的结论。')
    d.add_page_break()
    heading(d,'四 演示案例与交付使用')
    para(d,'演示采用独立生成的模拟零件清单，参考表有 7 个编号条目，待排序表有 8 行。实际运行后匹配 7 行、未匹配 1 行，其中 1 行关联两个参考候选。两条 M-002 保持原有先后关系；M-099 位于末尾。新 Excel 与逐行审计文件由录制过程中的程序直接生成。')
    image=ROOT/'docs/assets/实际结果.png'
    if image.exists():
        d.add_picture(str(image),width=Cm(16.3))
        p=para(d,'图 1  桌面程序本轮模拟案例的实际排序结果');p.paragraph_format.space_after=Pt(14)
    heading(d,'运行与复现')
    para(d,'Windows 用户可在完整源码 ZIP 中打开 Windows 目录，运行 Xuheng.exe；也可安装 Python 3.10 及以上版本，按依赖清单安装后运行 main.py。点击“填入模拟样例”“导入并识别”“开始对齐”，核对结果后填写保存位置并导出。')
    para(d,'源码目录包含安装启动说明、算法架构说明、输入样例、真实运行汇总和测试脚本。demo 目录提供解说词、SRT 字幕、视频实测记录及本轮导出文件；validation 目录保存附件汇总、千行计时和循环运行记录。')
    para(d,'项目的实际作用是把重复的报表顺序整理做成可复核的操作。使用者能够看到编号关联和行号变化，保留原资料，也保留对多候选项的人工核对依据。')
    for border in list(d._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    path.parent.mkdir(parents=True,exist_ok=True);d.save(path)
    print(path)

def main():
    p=argparse.ArgumentParser();p.add_argument('--private-info');args=p.parse_args()
    build(ROOT/'docs/项目申报书公开版.docx',{'name':'公开版已隐藏','organization':'公开版已隐藏','contact':'公开版已隐藏'})
    if args.private_info:
        build(ROOT/'delivery/项目申报书.docx',json.loads(Path(args.private_info).read_text(encoding='utf-8')))

if __name__=='__main__':main()
